"""
Exam Grouping core logic.
Extracted from the original script so it can be reused by the web app.
"""

from __future__ import annotations

import os
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

import networkx as nx
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# -----------------------------------------------------------------------------
# TEXT / COLUMN HELPERS
# -----------------------------------------------------------------------------
def clean_text(value) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    return text


def normalise_reg_no(value) -> str:
    """Keep registration numbers as text and remove Excel's trailing .0."""
    text = clean_text(value)
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return text.strip()


def detect_type(*values: object) -> Tuple[str, str]:
    """
    Returns (type, warning).
    User's rule:
      (P) => Practical
      (T) => Theory
      neither => Theory
    If both markers occur in the same combined record, Practical is chosen
    conservatively and a warning is returned for manual checking.
    """
    text = " ".join(clean_text(v).upper() for v in values)
    has_t = bool(re.search(r"\(\s*T\s*\)", text))
    has_p = bool(re.search(r"\(\s*P\s*\)", text))
    if has_t and has_p:
        return "P", "Both (T) and (P) found; classified as Practical for safety."
    if has_p:
        return "P", ""
    return "T", ""  # (T), or no marker at all


def remove_type_markers(text: object) -> str:
    value = clean_text(text)
    value = re.sub(r"\s*\(\s*[TP]\s*\)", "", value, flags=re.I)
    value = re.sub(r"\s+", " ", value)
    return value.strip(" :-")


def find_column(columns: Iterable[object], candidates: Sequence[str]):
    lookup = {str(c).strip().lower(): c for c in columns}
    for candidate in candidates:
        if candidate.lower() in lookup:
            return lookup[candidate.lower()]
    return None


def find_reg_column(columns: Iterable[object]):
    return find_column(
        columns,
        [
            "Reg No.", "Reg No", "Register No.", "Register No",
            "Registration No.", "Registration No", "Register Number",
            "Registration Number", "RegNo", "RegisterNo",
        ],
    )


def is_reg_header(value: object) -> bool:
    return clean_text(value).lower() in {
        "reg no.", "reg no", "register no.", "register no",
        "registration no.", "registration no", "register number",
        "registration number", "regno", "registerno",
    }


# -----------------------------------------------------------------------------
# SUBJECT PARSER
# -----------------------------------------------------------------------------
COURSE_CODE_PATTERN = re.compile(
    r"\b([A-Za-z][A-Za-z0-9._/-]*\d[A-Za-z0-9._/-]*(?:\s*\(\s*[TP]\s*\))?)\b",
    flags=re.I,
)


def parse_subject_name(subject: object) -> Tuple[str, str, str]:
    """
    Extract (clean_code, clean_name, raw_code) from strings such as:
      Introduction to Algebraic Structures ( SBU24MA3DSC201)
      Wave Optics(T) ( SBU24PH3DSC202(T))
      Properties of Solids and Fluids (P) ( SBU24PH3DSC201(P))
      Python Programming ( SBU24CS3SEC200(T))
    """
    text = clean_text(subject)
    matches = list(COURSE_CODE_PATTERN.finditer(text))
    if not matches:
        return "", remove_type_markers(text), ""
    # In the supplied layouts the course code is the last code-like token.
    match = matches[-1]
    raw_code = match.group(1).strip()
    code = remove_type_markers(raw_code)
    name_part = text[: match.start()].strip()
    # Remove the opening '(' that encloses the course code.
    name_part = re.sub(r"\s*\(\s*$", "", name_part).strip()
    name = remove_type_markers(name_part)
    return code, name, raw_code


# -----------------------------------------------------------------------------
# INPUT FORMAT DETECTION + STANDARDISATION
# -----------------------------------------------------------------------------
def locate_table_header(raw: pd.DataFrame, max_rows: int = 40):
    """Find a normal table header row, if present."""
    for i in range(min(max_rows, len(raw))):
        vals = [clean_text(v).lower() for v in raw.iloc[i].tolist()]
        has_reg = any(
            v in vals for v in {
                "reg no.", "reg no", "register no.", "register no",
                "registration no.", "registration no", "register number",
                "registration number",
            }
        )
        has_course = any(
            v in vals for v in {"course code", "course name", "subject name", "paper name"}
        )
        if has_reg and has_course:
            return i
    return None


def process_table_sheet(
    path: str,
    category: str,
    sheet_name: str,
    header_row: int,
) -> Tuple[List[dict], List[dict]]:
    df = pd.read_excel(path, sheet_name=sheet_name, header=header_row, dtype=object)
    reg_col = find_reg_column(df.columns)
    code_col = find_column(df.columns, ["Course Code"])
    name_col = find_column(df.columns, ["Course Name"])
    subject_col = find_column(df.columns, ["Subject Name"])
    paper_col = find_column(df.columns, ["Paper Name"])
    records: List[dict] = []
    warnings: List[dict] = []
    if reg_col is None:
        warnings.append({
            "Source File": os.path.basename(path),
            "Sheet": sheet_name,
            "Reg No.": "",
            "Warning": "Registration number column not found.",
        })
        return records, warnings
    for _, row in df.iterrows():
        reg = normalise_reg_no(row.get(reg_col, ""))
        if not reg:
            continue
        raw_code = clean_text(row.get(code_col, "")) if code_col is not None else ""
        raw_name = clean_text(row.get(name_col, "")) if name_col is not None else ""
        raw_subject = clean_text(row.get(subject_col, "")) if subject_col is not None else ""
        raw_paper = clean_text(row.get(paper_col, "")) if paper_col is not None else ""
        # Subject Name can fill missing code/name.
        if raw_subject and (not raw_code or not raw_name):
            parsed_code, parsed_name, _ = parse_subject_name(raw_subject)
            if not raw_code:
                raw_code = parsed_code
            if not raw_name:
                raw_name = parsed_name
        # Paper Name is a final name fallback.
        if not raw_name and raw_paper:
            raw_name = raw_paper
        course_type, type_warning = detect_type(
            raw_code, raw_name, raw_subject, raw_paper
        )
        course_code = remove_type_markers(raw_code)
        course_name = remove_type_markers(raw_name)
        if not course_code:
            # If code is absent but Paper Name itself contains a parseable code.
            p_code, p_name, _ = parse_subject_name(raw_paper)
            if p_code:
                course_code = p_code
                if not course_name:
                    course_name = p_name
        if not course_code:
            warnings.append({
                "Source File": os.path.basename(path),
                "Sheet": sheet_name,
                "Reg No.": reg,
                "Warning": "Course Code could not be identified; record skipped.",
            })
            continue
        if not course_name:
            course_name = course_code
            warnings.append({
                "Source File": os.path.basename(path),
                "Sheet": sheet_name,
                "Reg No.": reg,
                "Warning": f"Course Name missing for {course_code}; code used as name.",
            })
        if type_warning:
            warnings.append({
                "Source File": os.path.basename(path),
                "Sheet": sheet_name,
                "Reg No.": reg,
                "Warning": f"{course_code}: {type_warning}",
            })
        records.append({
            "Source File": os.path.basename(path),
            "Student Category": category,
            "Sheet": sheet_name,
            "Reg No.": reg,
            "Course Code": course_code,
            "Course Name": course_name,
            "Type": course_type,
        })
    return records, warnings


def process_subject_section_sheet(
    path: str,
    category: str,
    sheet_name: str,
) -> Tuple[List[dict], List[dict]]:
    """
    Processes sheets in which each course starts with:
        Subject Name : Course Name (Course Code)
    followed by a student table containing Register No.
    """
    raw = pd.read_excel(path, sheet_name=sheet_name, header=None, dtype=object)
    records: List[dict] = []
    warnings: List[dict] = []
    current_code = ""
    current_name = ""
    current_type = "T"
    current_subject_text = ""
    reg_col_index = None
    for row_index in range(len(raw)):
        values = [clean_text(v) for v in raw.iloc[row_index].tolist()]
        joined = " ".join(v for v in values if v)
        # Detect a new subject section.
        subject_match = re.search(r"Subject\s*Name\s*:\s*(.*)", joined, flags=re.I)
        if subject_match:
            current_subject_text = subject_match.group(1).strip()
            current_code, current_name, raw_code = parse_subject_name(current_subject_text)
            current_type, type_warning = detect_type(current_subject_text, raw_code)
            reg_col_index = None
            if not current_code:
                warnings.append({
                    "Source File": os.path.basename(path),
                    "Sheet": sheet_name,
                    "Reg No.": "",
                    "Warning": f"Could not extract Course Code from Subject Name: {current_subject_text}",
                })
            if not current_name and current_code:
                current_name = current_code
            if type_warning:
                warnings.append({
                    "Source File": os.path.basename(path),
                    "Sheet": sheet_name,
                    "Reg No.": "",
                    "Warning": f"{current_code or current_subject_text}: {type_warning}",
                })
            continue
        # Ignore anything before the first Subject Name section.
        if not current_code:
            continue
        # Locate the Register No column for this section.
        if reg_col_index is None:
            for idx, value in enumerate(values):
                if is_reg_header(value):
                    reg_col_index = idx
                    break
            continue
        # Read the registration number only from the identified Register No column.
        if reg_col_index < len(values):
            reg = normalise_reg_no(values[reg_col_index])
        else:
            reg = ""
        # Typical registration numbers contain digits; skip headings/totals/blanks.
        if not reg or not re.search(r"\d", reg):
            continue
        # Avoid accidentally treating Sl No / totals / labels as a registration number.
        if reg.lower() in {"total", "nil", "none"}:
            continue
        records.append({
            "Source File": os.path.basename(path),
            "Student Category": category,
            "Sheet": sheet_name,
            "Reg No.": reg,
            "Course Code": current_code,
            "Course Name": current_name,
            "Type": current_type,
        })
    return records, warnings


def process_workbook(path: str, category: str) -> Tuple[List[dict], List[dict]]:
    records: List[dict] = []
    warnings: List[dict] = []
    xls = pd.ExcelFile(path)
    for sheet_name in xls.sheet_names:
        raw = pd.read_excel(path, sheet_name=sheet_name, header=None, dtype=object)
        header_row = locate_table_header(raw)
        if header_row is not None:
            r, w = process_table_sheet(path, category, sheet_name, header_row)
        else:
            r, w = process_subject_section_sheet(path, category, sheet_name)
        records.extend(r)
        warnings.extend(w)
    if not records:
        warnings.append({
            "Source File": os.path.basename(path),
            "Sheet": "",
            "Reg No.": "",
            "Warning": "No usable student-course records were found in this workbook.",
        })
    return records, warnings


# -----------------------------------------------------------------------------
# CANONICAL COURSE NAMES / DUPLICATES
# -----------------------------------------------------------------------------
def apply_canonical_names(df: pd.DataFrame, warnings: List[dict]) -> pd.DataFrame:
    """For the same Course Code + Type, use the most frequent nonblank name."""
    df = df.copy()
    df["Paper ID"] = df["Course Code"].str.strip() + "-" + df["Type"].str.strip()
    canonical: Dict[str, str] = {}
    for paper_id, group in df.groupby("Paper ID", sort=False):
        names = [n.strip() for n in group["Course Name"].astype(str) if n.strip()]
        if not names:
            canonical_name = group.iloc[0]["Course Code"]
        else:
            canonical_name = Counter(names).most_common(1)[0][0]
        canonical[paper_id] = canonical_name
        unique_names = sorted(set(names), key=str.lower)
        if len(unique_names) > 1:
            warnings.append({
                "Source File": "ALL FILES",
                "Sheet": "",
                "Reg No.": "",
                "Warning": (
                    f"Different Course Names found for {paper_id}: "
                    + " | ".join(unique_names)
                    + f". Using: {canonical_name}"
                ),
            })
    df["Course Name"] = df["Paper ID"].map(canonical)
    return df


def deduplicate_student_papers(df: pd.DataFrame, warnings: List[dict]) -> pd.DataFrame:
    """
    One student-paper registration is enough for conflict analysis.
    If the same registration occurs in several source files, combine its metadata.
    """
    before = len(df)
    grouped_rows = []
    for (reg, paper_id), group in df.groupby(["Reg No.", "Paper ID"], sort=False):
        first = group.iloc[0]
        grouped_rows.append({
            "Source File": "; ".join(sorted(set(group["Source File"].astype(str)))),
            "Student Category": "; ".join(sorted(set(group["Student Category"].astype(str)))),
            "Sheet": "; ".join(sorted(set(group["Sheet"].astype(str)))),
            "Reg No.": reg,
            "Course Code": first["Course Code"],
            "Course Name": first["Course Name"],
            "Type": first["Type"],
            "Paper ID": paper_id,
        })
    dedup = pd.DataFrame(grouped_rows)
    removed = before - len(dedup)
    if removed:
        warnings.append({
            "Source File": "ALL FILES",
            "Sheet": "",
            "Reg No.": "",
            "Warning": f"{removed} duplicate student-paper row(s) combined across the input files.",
        })
    return dedup


# -----------------------------------------------------------------------------
# CONFLICT GRAPH + GROUPING
# -----------------------------------------------------------------------------
def build_conflict_graph(type_df: pd.DataFrame) -> Tuple[nx.Graph, pd.DataFrame]:
    """Create one node per paper and an edge when at least one student is common."""
    graph = nx.Graph()
    paper_info = (
        type_df[["Paper ID", "Course Code", "Course Name", "Type"]]
        .drop_duplicates("Paper ID")
        .set_index("Paper ID")
    )
    graph.add_nodes_from(paper_info.index.tolist())
    # Student -> papers is much faster than comparing every pair of columns.
    pair_students: Dict[Tuple[str, str], set] = defaultdict(set)
    for reg, group in type_df.groupby("Reg No."):
        papers = sorted(set(group["Paper ID"]))
        for i in range(len(papers)):
            for j in range(i + 1, len(papers)):
                pair_students[(papers[i], papers[j])].add(reg)
    conflict_rows = []
    for (p1, p2), students in pair_students.items():
        graph.add_edge(p1, p2, common_students=len(students))
        i1 = paper_info.loc[p1]
        i2 = paper_info.loc[p2]
        sorted_students = sorted(students)
        conflict_rows.append({
            "Paper 1 Code": i1["Course Code"],
            "Paper 1 Name": i1["Course Name"],
            "Paper 2 Code": i2["Course Code"],
            "Paper 2 Name": i2["Course Name"],
            "Common Student Count": len(sorted_students),
            "Common Reg Nos.": ", ".join(sorted_students),
        })
    conflicts = pd.DataFrame(conflict_rows)
    if not conflicts.empty:
        conflicts = conflicts.sort_values(
            ["Common Student Count", "Paper 1 Code", "Paper 2 Code"],
            ascending=[False, True, True],
        ).reset_index(drop=True)
    return graph, conflicts


def best_greedy_coloring(graph: nx.Graph) -> Tuple[Dict[str, int], str]:
    """
    Try several NetworkX greedy strategies and keep the solution using the
    fewest colors. This produces a strong practical timetable grouping, though
    like all greedy graph-coloring methods it is not a proof of the chromatic
    minimum for arbitrary graphs.
    """
    if graph.number_of_nodes() == 0:
        return {}, "none"
    strategies = [
        "saturation_largest_first",  # DSATUR
        "largest_first",
        "smallest_last",
        "connected_sequential_bfs",
        "connected_sequential_dfs",
    ]
    best = None
    best_strategy = ""
    best_count = None
    for strategy in strategies:
        try:
            coloring = nx.coloring.greedy_color(graph, strategy=strategy)
            color_count = (max(coloring.values()) + 1) if coloring else 0
            if best is None or color_count < best_count:
                best = coloring
                best_strategy = strategy
                best_count = color_count
        except Exception:
            continue
    if best is None:
        best = nx.coloring.greedy_color(graph, strategy="largest_first")
        best_strategy = "largest_first"
    return best, best_strategy


def make_group_rows(
    type_df: pd.DataFrame,
    graph: nx.Graph,
    coloring: Dict[str, int],
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Return final-output-style rows and a compact group summary."""
    if type_df.empty:
        return (
            pd.DataFrame(columns=["Course Code", "Course Name", "Students"]),
            pd.DataFrame(columns=["Group", "No. of Papers", "Total Student Registrations"]),
        )
    paper_info = (
        type_df.groupby(["Paper ID", "Course Code", "Course Name"], as_index=False)
        .agg(Students=("Reg No.", "nunique"))
        .set_index("Paper ID")
    )
    groups: Dict[int, List[str]] = defaultdict(list)
    for paper_id, color in coloring.items():
        groups[color].append(paper_id)
    # Order groups by color; within each group put larger/high-conflict papers first.
    output_rows = []
    summary_rows = []
    for group_no, color in enumerate(sorted(groups), start=1):
        papers = groups[color]
        papers = sorted(
            papers,
            key=lambda p: (-graph.degree[p], -int(paper_info.loc[p, "Students"]), str(p)),
        )
        output_rows.append({"Course Code": f"Group {group_no}", "Course Name": "", "Students": ""})
        student_registrations = 0
        for paper_id in papers:
            info = paper_info.loc[paper_id]
            student_count = int(info["Students"])
            student_registrations += student_count
            output_rows.append({
                "Course Code": info["Course Code"],
                "Course Name": info["Course Name"],
                "Students": student_count,
            })
        output_rows.append({"Course Code": "", "Course Name": "", "Students": ""})
        summary_rows.append({
            "Group": f"Group {group_no}",
            "No. of Papers": len(papers),
            "Total Student Registrations": student_registrations,
        })
    return pd.DataFrame(output_rows), pd.DataFrame(summary_rows)


def verify_grouping(graph: nx.Graph, coloring: Dict[str, int]) -> List[Tuple[str, str]]:
    """Return any conflicting paper pairs that accidentally received the same group."""
    errors = []
    for u, v in graph.edges():
        if coloring.get(u) == coloring.get(v):
            errors.append((u, v))
    return errors


# -----------------------------------------------------------------------------
# SUMMARY TABLES
# -----------------------------------------------------------------------------
def create_course_summary(raw_df: pd.DataFrame, dedup_df: pd.DataFrame) -> pd.DataFrame:
    if dedup_df.empty:
        return pd.DataFrame()
    rows = []
    for paper_id, group in dedup_df.groupby("Paper ID"):
        first = group.iloc[0]
        # Category counts use the original standardized rows so a student's source
        # category remains visible even if duplicate rows were later combined.
        raw_group = raw_df[raw_df["Paper ID"] == paper_id]
        regular = raw_group.loc[
            raw_group["Student Category"].str.contains("Regular", case=False, na=False),
            "Reg No.",
        ].nunique()
        supp = raw_group.loc[
            raw_group["Student Category"].str.contains("Improvement|Supplementary", case=False, na=False, regex=True),
            "Reg No.",
        ].nunique()
        rows.append({
            "Course Code": first["Course Code"],
            "Course Name": first["Course Name"],
            "Type": "Theory" if first["Type"] == "T" else "Practical",
            "Regular Students": int(regular),
            "Improvement/Supplementary Students": int(supp),
            "Total Unique Students": int(group["Reg No."].nunique()),
            "Source Files": "; ".join(sorted(set(raw_group["Source File"].astype(str)))),
        })
    return pd.DataFrame(rows).sort_values(
        ["Type", "Course Code", "Course Name"],
        ignore_index=True,
    )


# -----------------------------------------------------------------------------
# EXCEL OUTPUT + FORMATTING
# -----------------------------------------------------------------------------
def style_output_workbook(output_path: str) -> None:
    wb = load_workbook(output_path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    group_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(color="FFFFFF", bold=True)
    group_font = Font(bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        # Header row
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        # Group header rows in Theory/Practical output.
        if ws.title in {"Theory Groups", "Practical Groups"}:
            ws.auto_filter.ref = None  # group-header/blank-row layout is not ideal for filtering
            for row in ws.iter_rows(min_row=2):
                value = clean_text(row[0].value)
                if re.fullmatch(r"Group\s+\d+", value, flags=re.I):
                    for cell in row:
                        cell.fill = group_fill
                        cell.font = group_font
        # Wrap and sensible widths.
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        widths = {}
        for row in ws.iter_rows():
            for cell in row:
                value = "" if cell.value is None else str(cell.value)
                widths[cell.column] = min(max(widths.get(cell.column, 0), len(value) + 2), 45)
        for col_idx, width in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = max(10, width)
        ws.row_dimensions[1].height = 30
    wb.save(output_path)


# -----------------------------------------------------------------------------
# MAIN GENERATOR
# -----------------------------------------------------------------------------
def generate_exam_groups(
    regular_files: Sequence[str],
    supplementary_files: Sequence[str],
    output_path: str,
) -> dict:
    file_specs: List[Tuple[str, str]] = []
    file_specs.extend((f, "Regular") for f in regular_files)
    file_specs.extend((f, "Improvement/Supplementary") for f in supplementary_files)
    if not file_specs:
        raise ValueError("No input files selected.")
    records: List[dict] = []
    warnings: List[dict] = []
    for path, category in file_specs:
        if not os.path.exists(path):
            warnings.append({
                "Source File": os.path.basename(path),
                "Sheet": "",
                "Reg No.": "",
                "Warning": "Input file not found.",
            })
            continue
        r, w = process_workbook(path, category)
        records.extend(r)
        warnings.extend(w)
    if not records:
        raise ValueError("No usable student-course records were found in the selected files.")
    raw_df = pd.DataFrame(records)
    raw_df = apply_canonical_names(raw_df, warnings)
    # Preserve all standardized source rows for audit/category counts.
    standardized_all = raw_df.copy()
    # Deduplicate student + paper for actual conflict analysis.
    dedup_df = deduplicate_student_papers(raw_df, warnings)
    theory_df = dedup_df[dedup_df["Type"] == "T"].copy()
    practical_df = dedup_df[dedup_df["Type"] == "P"].copy()
    theory_graph, theory_conflicts = build_conflict_graph(theory_df)
    practical_graph, practical_conflicts = build_conflict_graph(practical_df)
    theory_coloring, theory_strategy = best_greedy_coloring(theory_graph)
    practical_coloring, practical_strategy = best_greedy_coloring(practical_graph)
    theory_errors = verify_grouping(theory_graph, theory_coloring)
    practical_errors = verify_grouping(practical_graph, practical_coloring)
    if theory_errors:
        warnings.append({
            "Source File": "PROGRAM",
            "Sheet": "Theory Groups",
            "Reg No.": "",
            "Warning": f"INTERNAL CHECK FAILED: {len(theory_errors)} theory grouping conflict(s).",
        })
    if practical_errors:
        warnings.append({
            "Source File": "PROGRAM",
            "Sheet": "Practical Groups",
            "Reg No.": "",
            "Warning": f"INTERNAL CHECK FAILED: {len(practical_errors)} practical grouping conflict(s).",
        })
    theory_groups, theory_group_summary = make_group_rows(
        theory_df, theory_graph, theory_coloring
    )
    practical_groups, practical_group_summary = make_group_rows(
        practical_df, practical_graph, practical_coloring
    )
    if theory_conflicts.empty:
        theory_conflicts = pd.DataFrame(columns=[
            "Paper 1 Code", "Paper 1 Name", "Paper 2 Code", "Paper 2 Name",
            "Common Student Count", "Common Reg Nos."
        ])
    theory_conflicts.insert(0, "Type", "Theory")
    if practical_conflicts.empty:
        practical_conflicts = pd.DataFrame(columns=[
            "Paper 1 Code", "Paper 1 Name", "Paper 2 Code", "Paper 2 Name",
            "Common Student Count", "Common Reg Nos."
        ])
    practical_conflicts.insert(0, "Type", "Practical")
    conflicts = pd.concat([theory_conflicts, practical_conflicts], ignore_index=True)
    course_summary = create_course_summary(standardized_all, dedup_df)
    group_summary = pd.concat([
        theory_group_summary.assign(Type="Theory"),
        practical_group_summary.assign(Type="Practical"),
    ], ignore_index=True)
    if not group_summary.empty:
        group_summary = group_summary[["Type", "Group", "No. of Papers", "Total Student Registrations"]]
    warnings_df = pd.DataFrame(
        warnings,
        columns=["Source File", "Sheet", "Reg No.", "Warning"],
    )
    # Add an overview / audit sheet.
    overview = pd.DataFrame([
        ["Input files", len(file_specs)],
        ["Regular files", len(regular_files)],
        ["Improvement/Supplementary files", len(supplementary_files)],
        ["Standardized source rows", len(standardized_all)],
        ["Unique student-paper registrations", len(dedup_df)],
        ["Unique students", dedup_df["Reg No."].nunique()],
        ["Theory papers", theory_df["Paper ID"].nunique()],
        ["Theory groups", (max(theory_coloring.values()) + 1) if theory_coloring else 0],
        ["Theory coloring strategy", theory_strategy],
        ["Practical papers", practical_df["Paper ID"].nunique()],
        ["Practical groups", (max(practical_coloring.values()) + 1) if practical_coloring else 0],
        ["Practical coloring strategy", practical_strategy],
        ["Theory grouping verification", "PASS" if not theory_errors else "FAIL"],
        ["Practical grouping verification", "PASS" if not practical_errors else "FAIL"],
        ["Warnings", len(warnings_df)],
    ], columns=["Item", "Value"])
    # Order standardized columns for checking.
    standardized_all = standardized_all[[
        "Source File", "Student Category", "Sheet", "Reg No.",
        "Course Code", "Course Name", "Type", "Paper ID"
    ]].sort_values(["Type", "Course Code", "Reg No."], ignore_index=True)
    dedup_out = dedup_df[[
        "Source File", "Student Category", "Sheet", "Reg No.",
        "Course Code", "Course Name", "Type", "Paper ID"
    ]].sort_values(["Type", "Course Code", "Reg No."], ignore_index=True)
    output_path = str(Path(output_path))
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        overview.to_excel(writer, sheet_name="Overview", index=False)
        theory_groups.to_excel(writer, sheet_name="Theory Groups", index=False)
        practical_groups.to_excel(writer, sheet_name="Practical Groups", index=False)
        group_summary.to_excel(writer, sheet_name="Group Summary", index=False)
        course_summary.to_excel(writer, sheet_name="Course Summary", index=False)
        conflicts.to_excel(writer, sheet_name="Conflicts", index=False)
        dedup_out.to_excel(writer, sheet_name="Combined Students", index=False)
        standardized_all.to_excel(writer, sheet_name="Standardized Data", index=False)
        warnings_df.to_excel(writer, sheet_name="Warnings", index=False)
    style_output_workbook(output_path)
    return {
        "output": output_path,
        "standardized_rows": len(standardized_all),
        "combined_rows": len(dedup_df),
        "unique_students": int(dedup_df["Reg No."].nunique()),
        "theory_papers": int(theory_df["Paper ID"].nunique()),
        "theory_groups": int((max(theory_coloring.values()) + 1) if theory_coloring else 0),
        "practical_papers": int(practical_df["Paper ID"].nunique()),
        "practical_groups": int((max(practical_coloring.values()) + 1) if practical_coloring else 0),
        "warnings": len(warnings_df),
        "overview": overview,
        "group_summary": group_summary,
        "course_summary": course_summary,
        "warnings_df": warnings_df,
        "theory_groups_df": theory_groups,
        "practical_groups_df": practical_groups,
    }
